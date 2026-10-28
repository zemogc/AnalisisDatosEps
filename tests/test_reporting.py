import pytest
from openpyxl import load_workbook

from eps_analytics import database
from eps_analytics import formatting as fmt
from eps_analytics.analysis import ejecutar_analisis
from eps_analytics.cli import main
from eps_analytics.reporting import generar_excel, generar_graficos
from eps_analytics.reporting.charts import GRAFICOS


@pytest.fixture(scope="module")
def resultado(engine):
    return ejecutar_analisis(database.cargar_dataset(engine), database.cargar_medicamentos(engine))


def test_excel_tiene_todas_las_hojas(resultado, tmp_path):
    ruta = generar_excel(resultado, tmp_path / "reporte.xlsx")
    libro = load_workbook(ruta)
    assert libro.sheetnames[0] == "Resumen"
    assert {"Por EPS", "Por Diagnóstico", "Tendencia Mensual", "Detalle"} <= set(libro.sheetnames)
    assert len(libro["Resumen"]._charts) == 2
    assert libro["Detalle"].max_row == resultado.kpis["consultas"] + 1


def test_excel_formato_moneda(resultado, tmp_path):
    libro = load_workbook(generar_excel(resultado, tmp_path / "reporte.xlsx"))
    hoja = libro["Por EPS"]
    encabezados = [c.value for c in hoja[1]]
    col = encabezados.index("Costo_Total") + 1
    assert "$" in hoja.cell(row=2, column=col).number_format


def test_graficos_se_generan(resultado, tmp_path):
    rutas = generar_graficos(resultado, tmp_path)
    assert {r.name for r in rutas} == set(GRAFICOS)
    assert all(r.stat().st_size > 10_000 for r in rutas)


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        (1_234_567, "$ 1.234.567"),
        (0, "$ 0"),
    ],
)
def test_formato_cop(valor, esperado):
    assert fmt.cop(valor) == esperado


@pytest.mark.parametrize(
    ("valor", "esperado"),
    [
        (1_720_000_000, "$ 1,72 MM"),
        (85_300_000, "$ 85,3 M"),
        (95_400, "$ 95 mil"),
        (500, "$ 500"),
    ],
)
def test_formato_cop_compacto(valor, esperado):
    assert fmt.cop_compacto(valor) == esperado


def test_formato_porcentaje():
    assert fmt.porcentaje(0.2534) == "25,3 %"


def test_cli_seed_y_report(tmp_path, monkeypatch):
    monkeypatch.setenv("EPS_DB_URL", f"sqlite:///{(tmp_path / 'cli.db').as_posix()}")
    monkeypatch.setenv("EPS_OUTPUT_DIR", str(tmp_path / "salida"))
    assert main(["seed", "--pacientes", "80", "--desde", "2024-01-01", "--hasta", "2024-06-30"]) == 0
    # Sin --reset no sobrescribe una base existente
    assert main(["seed", "--pacientes", "80"]) == 1
    assert main(["report"]) == 0
    assert (tmp_path / "salida" / "reporte_eps.xlsx").exists()
    assert len(list((tmp_path / "salida" / "graficos").glob("*.png"))) == len(GRAFICOS)


def test_cli_report_sin_base(tmp_path, monkeypatch):
    monkeypatch.setenv("EPS_DB_URL", f"sqlite:///{(tmp_path / 'vacia.db').as_posix()}")
    assert main(["report", "-o", str(tmp_path)]) == 1
