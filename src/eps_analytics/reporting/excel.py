"""Reporte Excel multi-hoja con formato y gráficos nativos de Excel."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from eps_analytics.analysis import ResultadoAnalisis

FORMATO_COP = '"$" #,##0'
FORMATO_PCT = "0.0%"
FORMATO_NUM = "#,##0"
FORMATO_DEC = "#,##0.00"

_RELLENO_ENCABEZADO = PatternFill("solid", fgColor="184F95")
_FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF")
_BORDE = Border(bottom=Side(style="thin", color="D9D9D9"))

_KPIS = [
    ("pacientes", "Pacientes atendidos", FORMATO_NUM),
    ("consultas", "Consultas registradas", FORMATO_NUM),
    ("costo_total", "Costo total", FORMATO_COP),
    ("costo_promedio_consulta", "Costo promedio por consulta", FORMATO_COP),
    ("costo_mediano_consulta", "Costo mediano por consulta", FORMATO_COP),
    ("costo_por_paciente", "Costo promedio por paciente", FORMATO_COP),
    ("consultas_por_paciente", "Consultas por paciente", FORMATO_DEC),
    ("participacion_medicamentos", "Participación de medicamentos en el costo", FORMATO_PCT),
]


def _formato_columna(nombre: str) -> str | None:
    n = nombre.lower()
    if "participacion" in n or "variacion" in n:
        return FORMATO_PCT
    if "costo" in n or "gasto" in n or "media_movil" in n:
        return FORMATO_COP
    if n in {"pacientes", "consultas", "formulaciones", "unidades", "num_medicamentos"}:
        return FORMATO_NUM
    return None


def _dar_formato(ws: Worksheet, df: pd.DataFrame, fila_encabezado: int = 1) -> None:
    for celda in ws[fila_encabezado]:
        celda.fill = _RELLENO_ENCABEZADO
        celda.font = _FUENTE_ENCABEZADO
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[fila_encabezado].height = 30
    ws.freeze_panes = ws.cell(row=fila_encabezado + 1, column=1)
    ws.auto_filter.ref = ws.dimensions

    for idx, celda in enumerate(ws[fila_encabezado], start=1):
        nombre = str(celda.value or "")
        formato = _formato_columna(nombre)
        letra = get_column_letter(idx)
        valores = [nombre] + [str(c.value) for c in ws[letra][fila_encabezado : fila_encabezado + 200]]
        ws.column_dimensions[letra].width = min(max(len(v) for v in valores) + 3, 45)
        if formato:
            for c in ws[letra][fila_encabezado:]:
                c.number_format = formato
                ws.column_dimensions[letra].width = max(ws.column_dimensions[letra].width, 16)
        if "fecha" in nombre.lower() or nombre == "Mes":
            for c in ws[letra][fila_encabezado:]:
                c.number_format = "yyyy-mm-dd" if nombre != "Mes" else "mmm yyyy"
            ws.column_dimensions[letra].width = 13


def _hoja(writer: pd.ExcelWriter, nombre: str, df: pd.DataFrame, index: bool = True) -> Worksheet:
    datos = df.reset_index() if index else df
    datos.to_excel(writer, sheet_name=nombre, index=False)
    ws = writer.sheets[nombre]
    _dar_formato(ws, datos)
    return ws


def _hoja_resumen(writer: pd.ExcelWriter, res: ResultadoAnalisis) -> None:
    ws = writer.book.create_sheet("Resumen", 0)
    ws["A1"] = "Reporte de costos de atención en salud por EPS"
    ws["A1"].font = Font(bold=True, size=16, color="184F95")
    periodo = res.datos["Fecha_Consulta"]
    ws["A2"] = (
        f"Periodo: {periodo.min():%Y-%m-%d} a {periodo.max():%Y-%m-%d}  ·  "
        f"Generado: {datetime.now():%Y-%m-%d %H:%M}"
    )
    ws["A2"].font = Font(italic=True, color="52514E")

    ws["A4"], ws["B4"] = "Indicador", "Valor"
    for c in ws[4][:2]:
        c.fill, c.font = _RELLENO_ENCABEZADO, _FUENTE_ENCABEZADO
    for i, (clave, etiqueta, formato) in enumerate(_KPIS, start=5):
        ws.cell(row=i, column=1, value=etiqueta).border = _BORDE
        celda = ws.cell(row=i, column=2, value=res.kpis[clave])
        celda.number_format, celda.border = formato, _BORDE
    fila = 5 + len(_KPIS)
    ws.cell(row=fila, column=1, value="Gasto concentrado en el 5 % de pacientes más costosos")
    ws.cell(row=fila, column=2, value=res.concentracion_top5).number_format = FORMATO_PCT
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 22

    hoja_eps = writer.sheets["Por EPS"]
    n = len(res.por_eps)
    barras = BarChart()
    barras.type = "bar"
    barras.title = "Costo total por EPS"
    barras.style = 10
    barras.add_data(Reference(hoja_eps, min_col=4, min_row=1, max_row=n + 1), titles_from_data=True)
    barras.set_categories(Reference(hoja_eps, min_col=1, min_row=2, max_row=n + 1))
    barras.legend = None
    barras.height, barras.width = 8, 16
    ws.add_chart(barras, "D4")

    hoja_mes = writer.sheets["Tendencia Mensual"]
    m = len(res.mensual)
    linea = LineChart()
    linea.title = "Costo total mensual"
    linea.style = 12
    linea.add_data(Reference(hoja_mes, min_col=3, min_row=1, max_row=m + 1), titles_from_data=True)
    linea.add_data(Reference(hoja_mes, min_col=6, min_row=1, max_row=m + 1), titles_from_data=True)
    linea.set_categories(Reference(hoja_mes, min_col=1, min_row=2, max_row=m + 1))
    linea.height, linea.width = 8, 16
    ws.add_chart(linea, "D21")


def generar_excel(res: ResultadoAnalisis, ruta: Path) -> Path:
    ruta.parent.mkdir(parents=True, exist_ok=True)
    detalle = res.datos.drop(columns=["Mes"]).sort_values("Fecha_Consulta")
    with pd.ExcelWriter(ruta, engine="openpyxl") as writer:
        _hoja(writer, "Por EPS", res.por_eps)
        _hoja(writer, "Por Diagnóstico", res.por_diagnostico)
        _hoja(writer, "Tendencia Mensual", res.mensual)
        _hoja(writer, "Grupo Etario", res.por_grupo_etario)
        _hoja(writer, "Tipo de Atención", res.por_tipo_atencion)
        _hoja(writer, "Matriz EPS x Dx", res.matriz_eps_dx)
        _hoja(writer, "Top Medicamentos", res.medicamentos)
        _hoja(writer, "Consultas Atípicas", res.atipicos, index=False)
        _hoja(writer, "Pacientes Alto Costo", res.alto_costo, index=False)
        _hoja(writer, "Detalle", detalle, index=False)
        _hoja_resumen(writer, res)
        # La matriz es de costos aunque sus columnas sean nombres de EPS
        ws = writer.sheets["Matriz EPS x Dx"]
        for fila in ws.iter_rows(min_row=2, min_col=2):
            for c in fila:
                c.number_format = FORMATO_COP
    return ruta
